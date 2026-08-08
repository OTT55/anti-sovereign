"""Multi-format document parsing. Each parser's only job is turning raw file bytes into
plain text — chunking, indexing, retrieval, and extraction don't know or care what the
source format was.
"""
import csv
import email
from email import policy
from io import BytesIO, StringIO


class UnsupportedFormat(Exception):
    """Raised for an unrecognized extension or a file that fails to parse."""


def parse_pdf(raw):
    import fitz  # PyMuPDF
    doc = fitz.open(stream=raw, filetype="pdf")
    try:
        return "\n\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def parse_docx(raw):
    import docx
    doc = docx.Document(BytesIO(raw))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def parse_xlsx(raw):
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(raw), data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_pptx(raw):
    import pptx
    prs = pptx.Presentation(BytesIO(raw))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        if texts:
            parts.append(f"# Slide {i}\n" + "\n".join(texts))
    return "\n\n".join(parts)


def parse_csv(raw):
    text = raw.decode("utf-8", errors="replace")
    rows = [", ".join(cell for cell in row if cell) for row in csv.reader(StringIO(text)) if any(row)]
    return "\n".join(rows)


def parse_eml(raw):
    msg = email.message_from_bytes(raw, policy=policy.default)
    header = (
        f"From: {msg.get('From', '')}\n"
        f"To: {msg.get('To', '')}\n"
        f"Subject: {msg.get('Subject', '')}\n"
        f"Date: {msg.get('Date', '')}\n"
    )
    body = msg.get_body(preferencelist=("plain", "html"))
    return header + "\n" + (body.get_content() if body else "")


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".pptx": parse_pptx,
    ".csv": parse_csv,
    ".eml": parse_eml,
}

SUPPORTED_EXTENSIONS = ".txt, .md, .pdf, .docx, .xlsx, .pptx, .csv, .eml"


def parse_file(filename, raw):
    """Return plain text extracted from `raw` bytes, or raise UnsupportedFormat."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in (".txt", ".md"):
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError:
            raise UnsupportedFormat(f"{filename} is not valid UTF-8 text.")
    parser = PARSERS.get(ext)
    if parser is None:
        raise UnsupportedFormat(
            f"'{ext or 'that file type'}' is not supported. Supported formats: {SUPPORTED_EXTENSIONS}."
        )
    try:
        text = parser(raw)
    except Exception as e:  # corrupt file, password-protected, malformed — surface honestly
        raise UnsupportedFormat(f"Could not parse {filename} ({type(e).__name__}: {e}).")
    if not text.strip():
        raise UnsupportedFormat(f"{filename} parsed successfully but contained no extractable text.")
    return text
